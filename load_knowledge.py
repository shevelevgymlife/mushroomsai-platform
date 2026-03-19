"""
Скрипт загрузки базы знаний из Google Drive в PostgreSQL.
Использование: python load_knowledge.py
"""

import os
import io
import sys
import tempfile
from pathlib import Path

# Принудительно переключаем вывод на UTF-8 для поддержки всех Unicode-символов
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
if hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")

from dotenv import load_dotenv
from googleapiclient.discovery import build
from googleapiclient.http import MediaIoBaseDownload
from google.oauth2 import service_account
import psycopg2

# Парсинг текста из файлов
try:
    import PyPDF2
except ImportError:
    PyPDF2 = None

try:
    from docx import Document as DocxDocument
except ImportError:
    DocxDocument = None

try:
    import openpyxl
except ImportError:
    openpyxl = None

load_dotenv()

DRIVE_FOLDER_ID = "1hT6KDF4V_4Dm1-Iz2diHx8ssCSI-cIiy"
SERVICE_ACCOUNT_FILE = "service_account.json"
SUPPORTED_MIME_TYPES = {
    "application/pdf": ".pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": ".docx",
    "text/plain": ".txt",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": ".xlsx",
    "application/msword": ".doc",
}

CREATE_KNOWLEDGE_BASE = """
CREATE TABLE IF NOT EXISTS knowledge_base (
    id SERIAL PRIMARY KEY,
    title TEXT,
    content TEXT,
    category TEXT,
    source_file TEXT,
    created_at TIMESTAMP DEFAULT NOW()
);
"""

CREATE_SHOP_PRODUCTS = """
CREATE TABLE IF NOT EXISTS shop_products (
    id SERIAL PRIMARY KEY,
    name TEXT,
    description TEXT,
    price INTEGER,
    url TEXT,
    mushroom_type TEXT,
    created_at TIMESTAMP DEFAULT NOW()
);
"""


def get_drive_service():
    if not Path(SERVICE_ACCOUNT_FILE).exists():
        print(f"[ERROR] Файл {SERVICE_ACCOUNT_FILE} не найден в текущей папке.")
        sys.exit(1)

    credentials = service_account.Credentials.from_service_account_file(
        SERVICE_ACCOUNT_FILE,
        scopes=["https://www.googleapis.com/auth/drive.readonly"],
    )
    return build("drive", "v3", credentials=credentials)


def list_files(service, folder_id):
    """Возвращает список файлов из папки Google Drive."""
    files = []
    page_token = None

    mime_filter = " or ".join(
        f"mimeType='{mime}'" for mime in SUPPORTED_MIME_TYPES
    )
    query = f"'{folder_id}' in parents and ({mime_filter}) and trashed=false"

    while True:
        response = service.files().list(
            q=query,
            fields="nextPageToken, files(id, name, mimeType)",
            pageToken=page_token,
            pageSize=100,
        ).execute()

        files.extend(response.get("files", []))
        page_token = response.get("nextPageToken")
        if not page_token:
            break

    return files


def download_file(service, file_id, file_name, mime_type):
    """Скачивает файл из Google Drive в временный файл, возвращает путь."""
    ext = SUPPORTED_MIME_TYPES.get(mime_type, "")
    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=ext)

    request = service.files().get_media(fileId=file_id)
    downloader = MediaIoBaseDownload(tmp, request)

    done = False
    while not done:
        _, done = downloader.next_chunk()

    tmp.close()
    return tmp.name


def extract_text_pdf(path):
    if PyPDF2 is None:
        return "[PyPDF2 не установлен — текст PDF недоступен]"
    text_parts = []
    with open(path, "rb") as f:
        reader = PyPDF2.PdfReader(f)
        for page in reader.pages:
            page_text = page.extract_text()
            if page_text:
                text_parts.append(page_text)
    return "\n".join(text_parts)


def extract_text_docx(path):
    if DocxDocument is None:
        return "[python-docx не установлен — текст DOCX недоступен]"
    doc = DocxDocument(path)
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())


def extract_text_txt(path):
    for encoding in ("utf-8", "cp1251", "latin-1"):
        try:
            with open(path, "r", encoding=encoding) as f:
                return f.read()
        except UnicodeDecodeError:
            continue
    return "[Не удалось декодировать текстовый файл]"


def extract_text_xlsx(path):
    if openpyxl is None:
        return "[openpyxl не установлен — текст XLSX недоступен]"
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = []
    for sheet in wb.worksheets:
        for row in sheet.iter_rows(values_only=True):
            row_text = "\t".join(str(cell) for cell in row if cell is not None)
            if row_text.strip():
                rows.append(row_text)
    return "\n".join(rows)


def extract_text(path, mime_type):
    if mime_type == "application/pdf":
        return extract_text_pdf(path)
    elif mime_type in (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "application/msword",
    ):
        return extract_text_docx(path)
    elif mime_type == "text/plain":
        return extract_text_txt(path)
    elif mime_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
        return extract_text_xlsx(path)
    return ""


def guess_category(file_name):
    name_lower = file_name.lower()
    if any(k in name_lower for k in ("рецепт", "recipe", "кулинар")):
        return "рецепты"
    if any(k in name_lower for k in ("протокол", "protocol", "схема", "лечение")):
        return "протоколы"
    if any(k in name_lower for k in ("исследован", "research", "наука", "science")):
        return "исследования"
    if any(k in name_lower for k in ("гриб", "mushroom", "чага", "рейши", "шиитаке", "кордицепс", "лев")):
        return "грибы"
    return "общее"


def main():
    database_url = os.getenv("DATABASE_URL")
    if not database_url:
        print("[ERROR] DATABASE_URL не задан в .env файле.")
        sys.exit(1)

    # Приводим asyncpg URL к обычному psycopg2
    pg_url = database_url.replace("postgresql+asyncpg://", "postgresql://")

    print("Подключение к Google Drive...")
    service = get_drive_service()

    print(f"Получение списка файлов из папки {DRIVE_FOLDER_ID}...")
    files = list_files(service, DRIVE_FOLDER_ID)

    if not files:
        print("Папка пуста или нет файлов поддерживаемых форматов.")
        return

    print(f"Найдено файлов: {len(files)}")

    print("Подключение к PostgreSQL...")
    conn = psycopg2.connect(pg_url)
    cur = conn.cursor()

    print("Создание таблиц...")
    cur.execute(CREATE_KNOWLEDGE_BASE)
    cur.execute(CREATE_SHOP_PRODUCTS)
    conn.commit()

    loaded = 0
    errors = 0

    for i, file in enumerate(files, 1):
        file_id = file["id"]
        file_name = file["name"]
        mime_type = file["mimeType"]

        print(f"[{i}/{len(files)}] Скачивание: {file_name} ...", end=" ")

        tmp_path = None
        try:
            tmp_path = download_file(service, file_id, file_name, mime_type)
            text = extract_text(tmp_path, mime_type)

            if not text.strip():
                print("пропущен (пустой текст)")
                continue

            category = guess_category(file_name)

            # Проверяем, не загружен ли уже этот файл
            cur.execute(
                "SELECT id FROM knowledge_base WHERE source_file = %s LIMIT 1",
                (file_name,),
            )
            existing = cur.fetchone()

            if existing:
                cur.execute(
                    "UPDATE knowledge_base SET title=%s, content=%s, category=%s WHERE source_file=%s",
                    (file_name, text, category, file_name),
                )
                print("обновлён")
            else:
                cur.execute(
                    "INSERT INTO knowledge_base (title, content, category, source_file) VALUES (%s, %s, %s, %s)",
                    (file_name, text, category, file_name),
                )
                print("загружен")

            conn.commit()
            loaded += 1

        except Exception as e:
            print(f"ОШИБКА: {e}")
            conn.rollback()
            errors += 1
        finally:
            if tmp_path and Path(tmp_path).exists():
                os.unlink(tmp_path)

    cur.close()
    conn.close()

    print(f"\nГотово! Загружено: {loaded}, ошибок: {errors}")
    print("Таблицы knowledge_base и shop_products готовы к использованию.")


if __name__ == "__main__":
    main()
